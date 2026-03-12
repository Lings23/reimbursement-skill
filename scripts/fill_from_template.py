#!/usr/bin/env python3
"""
使用模板文件自动化填写加班餐费报销审批表

用法:
    python fill_from_template.py --input 源文件.docx --output 输出文件.docx --data data.json
"""

import json
import argparse
from docx import Document
from pathlib import Path


class DocxFiller:
    def __init__(self, template_path):
        """加载模板配置"""
        with open(template_path, 'r', encoding='utf-8') as f:
            self.template = json.load(f)
        print(f"✓ 已加载模板：{self.template['template_name']} v{self.template['version']}")
    
    def fill(self, input_path, output_path, data):
        """填写文档"""
        doc = Document(input_path)
        table = doc.tables[self.template['document_structure']['table_index']]
        fields = self.template['fields']
        
        print(f"\n正在填写：{input_path}")
        
        # 填写用餐人数
        if '用餐人数' in data and '用餐人数' in fields:
            row = fields['用餐人数']['row']
            col = fields['用餐人数']['col']
            table.cell(row, col).text = str(data['用餐人数'])
            print(f"  ✓ 用餐人数：{data['用餐人数']}")
        
        # 填写用餐金额
        if '用餐金额' in data and '用餐金额' in fields:
            row = fields['用餐金额']['row']
            col = fields['用餐金额']['col']
            table.cell(row, col).text = str(data['用餐金额'])
            print(f"  ✓ 用餐金额：{data['用餐金额']} 元")
        
        # 填写用餐人员名单
        if '用餐人员名单' in data and '用餐人员名单' in fields:
            names = data['用餐人员名单']
            name_config = fields['用餐人员名单']
            
            # 按列分布姓名
            name_cols = name_config['name_columns']
            name_index = 0
            
            for col_config in name_cols:
                col = col_config['col']
                max_items = col_config['max_items']
                start_row = name_config['data_start_row']
                
                for i in range(max_items):
                    if name_index < len(names):
                        row = start_row + i
                        table.cell(row, col).text = names[name_index]
                        print(f"  ✓ 姓名 [{row},{col}]: {names[name_index]}")
                        name_index += 1
                    else:
                        break
        
        # 填写加班事由
        if '加班事由' in data and '加班事由' in fields:
            row = fields['加班事由']['row']
            col = fields['加班事由']['col']
            table.cell(row, col).text = data['加班事由']
            print(f"  ✓ 加班事由：{data['加班事由']}")
        
        # 保存
        doc.save(output_path)
        print(f"\n✓ 文档已保存：{output_path}")
        return output_path


def main():
    parser = argparse.ArgumentParser(description='使用模板填写加班餐费报销审批表')
    parser.add_argument('--input', '-i', required=True, help='输入 docx 文件')
    parser.add_argument('--output', '-o', required=True, help='输出 docx 文件')
    parser.add_argument('--template', '-t', default='templates/加班餐费报销审批表.json', 
                        help='模板 JSON 文件')
    parser.add_argument('--data', '-d', help='数据 JSON 文件')
    parser.add_argument('--names', '-n', nargs='+', help='用餐人员名单')
    parser.add_argument('--amount', '-a', type=float, help='用餐金额')
    
    args = parser.parse_args()
    
    # 加载数据
    data = {}
    if args.data:
        with open(args.data, 'r', encoding='utf-8') as f:
            data = json.load(f)
    if args.names:
        data['用餐人员名单'] = args.names
    if args.amount:
        data['用餐金额'] = args.amount
    
    # 自动计算用餐人数
    if '用餐人员名单' in data and '用餐人数' not in data:
        data['用餐人数'] = len(data['用餐人员名单'])
        print(f"自动计算用餐人数：{data['用餐人数']}")
    
    # 填写文档
    filler = DocxFiller(args.template)
    filler.fill(args.input, args.output, data)


if __name__ == '__main__':
    main()
